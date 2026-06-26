import uuid
import hashlib
from django.db import models
from django.conf import settings
import fitz
from docx import Document
import os


class DocumentoAporte(models.Model):
    ESTADOS = [
        ('pendiente', 'Pendiente de verificacion'),
        ('aceptado', 'Aceptado'),
        ('rechazado', 'Rechazado'),
    ]

    MOTIVOS_RECHAZO = [
        ('duplicado', 'Documento duplicado'),
        ('formato_incorrecto', 'Formato incorrecto'),
        ('contenido_no_academico', 'Contenido no academico'),
    ]

    usuario_interno = models.ForeignKey(
        settings.AUTH_USER_MODEL,
        on_delete=models.CASCADE,
        null=True, blank=True,
        related_name='aportes'
    )

    nombre_completo = models.CharField(max_length=200)
    email = models.EmailField(db_index=True)
    organizacion_origen = models.CharField(max_length=200, default='No especificado')

    titulo = models.CharField(max_length=300)
    descripcion = models.TextField()
    archivo = models.FileField(upload_to='aportes/%Y/%m/')
    archivo_hash = models.CharField(
        max_length=64, blank=True, db_index=True,
        help_text="Hash MD5 del archivo para detectar duplicados"
    )

    estado = models.CharField(
        max_length=20, choices=ESTADOS, default='pendiente'
    )

    motivo_rechazo = models.CharField(
        max_length=50,
        choices=MOTIVOS_RECHAZO,
        null=True,
        blank=True
    )
    comentario_rechazo = models.TextField(
        null=True,
        blank=True,
        help_text="Comentario adicional sobre el rechazo"
    )

    revisado_por = models.ForeignKey(
        settings.AUTH_USER_MODEL,
        null=True, blank=True,
        on_delete=models.SET_NULL,
        related_name='aportes_revisados'
    )
    revisado_at = models.DateTimeField(null=True, blank=True)

    created_at = models.DateTimeField(auto_now_add=True)

    class Meta:
        verbose_name = 'Documento Aporte'
        verbose_name_plural = 'Documentos Aporte'
        ordering = ['-created_at']

    def __str__(self):
        return f"{self.titulo} ({self.estado})"

    def calcular_hash(self):
        """Calcula hash MD5 del archivo para evitar duplicados."""
        if not self.archivo:
            return
        md5 = hashlib.md5()
        self.archivo.seek(0)
        for chunk in self.archivo.chunks():
            md5.update(chunk)
        self.archivo_hash = md5.hexdigest()
        self.archivo.seek(0)

    def save(self, *args, **kwargs):
        if self.archivo and not self.archivo_hash:
            self.calcular_hash()

        is_new = self.pk is None
        super().save(*args, **kwargs)

        if is_new and self.estado == 'pendiente':
            aprobado, motivo = self.verificar_automaticamente()
            if not aprobado:
                self.estado = 'rechazado'
                self.comentario_rechazo = motivo
                super().save(update_fields=['estado', 'comentario_rechazo'])

    def verificar_automaticamente(self):
        ext = os.path.splitext(self.archivo.name)[1].lower()
        if ext not in ['.pdf', '.docx']:
            return False, "Formato no permitido. Solo PDF o DOCX."

        if self.archivo.size < 50 * 1024:
            return False, "El archivo es demasiado pequeno. Parece estar vacio."

        if ext == '.pdf':
            try:
                doc = fitz.open(self.archivo.path)
                num_paginas = len(doc)

                if num_paginas < 5:
                    doc.close()
                    return False, f"El documento solo tiene {num_paginas} paginas. Minimo requerido: 5."

                texto = ""
                for pagina in doc:
                    texto += pagina.get_text()
                doc.close()

                palabras = len(texto.split())
                if palabras < 250:
                    return False, f"El documento tiene muy poco contenido ({palabras} palabras). Minimo: 250."

            except Exception as e:
                return False, f"No se pudo leer el PDF: {str(e)}"

        elif ext == '.docx':
            try:
                doc = Document(self.archivo.path)
                num_parrafos = len(doc.paragraphs)

                if num_parrafos < 10:
                    return False, f"El documento tiene muy pocos parrafos ({num_parrafos}). Minimo: 10."

                palabras = sum(len(p.text.split()) for p in doc.paragraphs if p.text)
                if palabras < 250:
                    return False, f"El documento tiene muy poco contenido ({palabras} palabras). Minimo: 250."

            except Exception as e:
                return False, f"No se pudo leer el DOCX: {str(e)}"

        self.calcular_hash()
        from biblioteca.models import TrabajoInvestigacion
        existe = TrabajoInvestigacion.objects.filter(
            archivo_ruta__icontains=self.archivo_hash[:16]
        ).exists()
        if existe:
            return False, "Este documento ya existe en nuestra biblioteca."

        return True, "Documento verificado automaticamente. Aprobado."

    def generar_creditos(self):
        """Genera 2 creditos de descarga. Solo externos reciben token."""
        from biblioteca.models import CreditoDescarga

        for i in range(2):
            token = None
            if not self.usuario_interno:
                token = str(uuid.uuid4()).replace('-', '')

            CreditoDescarga.objects.create(
                aporte=self,
                usuario=self.usuario_interno,
                email_externo=self.email if not self.usuario_interno else None,
                token_acceso=token
            )


class CreditoDescarga(models.Model):
    usuario = models.ForeignKey(
        settings.AUTH_USER_MODEL,
        on_delete=models.CASCADE,
        null=True, blank=True,
        related_name='creditos_descarga'
    )
    email_externo = models.EmailField(null=True, blank=True)
    token_acceso = models.CharField(
        max_length=64, null=True, blank=True,
        db_index=True, unique=True
    )

    aporte = models.ForeignKey(
        DocumentoAporte, on_delete=models.CASCADE,
        related_name='creditos'
    )
    usado = models.BooleanField(default=False)
    usado_en = models.DateTimeField(null=True, blank=True)
    usado_para = models.ForeignKey(
        'biblioteca.MaterialBibliografico',
        null=True, blank=True,
        on_delete=models.SET_NULL,
        related_name='creditos_descarga'
    )
    created_at = models.DateTimeField(auto_now_add=True)

    class Meta:
        verbose_name = 'Credito de Descarga'
        verbose_name_plural = 'Creditos de Descarga'

    def __str__(self):
        return f"Credito {self.id} ({'usado' if self.usado else 'disponible'})"


class AccesoExterno(models.Model):
    material = models.ForeignKey(
        'biblioteca.MaterialBibliografico',
        on_delete=models.CASCADE,
        null=True, blank=True,
        related_name='accesos_externos'
    )
    token_acceso = models.CharField(max_length=64, db_index=True)
    email = models.EmailField()
    nombre_completo = models.CharField(max_length=200)
    fecha_acceso = models.DateTimeField(auto_now_add=True)
    ip_origen = models.GenericIPAddressField(null=True, blank=True)

    class Meta:
        verbose_name = 'Acceso Externo'
        verbose_name_plural = 'Accesos Externos'
        unique_together = ['token_acceso', 'material']

    def __str__(self):
        return f"Acceso {self.email} -> {self.material.titulo if self.material else 'N/A'}"